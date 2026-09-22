"""Edexcel IGCSE Maths A (4MA1) Higher Tier — student topic RAG sheet."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

# ---- palette (matches the deck) ----
NAVY, AMBER_D, SLATE = "0D2B4E", "D4861B", "5A6B7C"
TINT, WHITE, HAIR = "EEF3F8", "FFFFFF", "DCE6F0"
G_BG, G_TX = "DCF0E4", "1B6B4A"
A_BG, A_TX = "FCEBCE", "8A5A00"
R_BG, R_TX = "FADCD9", "9B2C20"
F = "Calibri"

TOPICS = [
 ("1  Number and the number system", [
  "Place value, ordering, rounding to decimal places and significant figures",
  "Four operations with integers, decimals and negative numbers",
  "Factors, multiples, primes, HCF and LCM, prime factorisation",
  "Fractions: simplifying, comparing, all four operations, mixed numbers",
  "Converting between fractions, decimals and percentages",
  "Percentage of an amount; percentage increase and decrease",
  "Percentage change, profit and loss",
  "Reverse percentages (finding the original amount)",
  "Compound interest, depreciation, growth and decay",
  "Ratio: simplifying, sharing in a ratio, the unitary method",
  "Direct and inverse proportion",
  "Compound measures: speed, density and pressure",
  "Powers, roots and the index laws",
  "Negative and fractional indices",
  "Standard form: writing, converting and calculating",
  "Surds: simplifying and rationalising the denominator",
  "Upper and lower bounds; error intervals",
  "Recurring decimals into fractions",
  "Set language, set notation and Venn diagrams",
 ]),
 ("2  Equations, formulae and identities", [
  "Algebraic notation and collecting like terms",
  "Expanding single brackets",
  "Expanding double brackets",
  "Factorising using a common factor",
  "Factorising quadratics of the form x² + bx + c",
  "Factorising quadratics of the form ax² + bx + c",
  "The difference of two squares",
  "Substituting numbers into expressions and formulae",
  "Solving linear equations, including x on both sides",
  "Solving equations with brackets and with fractions",
  "Rearranging a formula to change the subject",
  "Rearranging where the new subject appears twice",
  "Simultaneous linear equations",
  "Simultaneous equations where one is quadratic",
  "Solving quadratic equations by factorising",
  "Solving quadratic equations with the formula",
  "Completing the square",
  "Algebraic fractions: simplifying and the four operations",
  "Solving equations that contain algebraic fractions",
  "Linear inequalities and representing them on a number line",
  "Quadratic inequalities",
  "Proportionality, including squares and cubes",
  "Forming an equation from a worded problem, then solving it",
  "Algebraic proof and reasoning",
 ]),
 ("3  Sequences, functions and graphs", [
  "Sequences: term-to-term rules and number patterns",
  "The nth term of a linear sequence",
  "The nth term of a quadratic sequence",
  "Coordinates and the midpoint of a line segment",
  "Plotting straight-line graphs",
  "Gradient, intercept and the equation y = mx + c",
  "Finding the equation of a line through two points",
  "Parallel and perpendicular lines",
  "Real-life graphs, including distance–time graphs",
  "Plotting and interpreting quadratic graphs",
  "Cubic and reciprocal graphs",
  "Exponential graphs, growth and decay",
  "Graphs of sine, cosine and tangent",
  "Solving equations graphically",
  "Function notation f(x)",
  "Composite and inverse functions",
  "Transformations of graphs",
  "Differentiating polynomials (the power rule)",
  "The gradient of a curve at a point, and tangents",
  "Turning points: maximum and minimum",
  "Kinematics: displacement, velocity and acceleration",
 ]),
 ("4  Geometry and trigonometry", [
  "Angle facts: on a line, at a point, and in a triangle",
  "Angles in parallel lines",
  "Interior and exterior angles of polygons",
  "Properties of triangles and quadrilaterals",
  "Congruence and similarity",
  "Similar shapes: length scale factor",
  "Similar shapes: area and volume scale factors",
  "Perimeter and area of rectangles, triangles, parallelograms, trapezia",
  "Circles: circumference and area",
  "Arcs, sectors and segments",
  "Circle theorems",
  "Volume and surface area of prisms and cylinders",
  "Volume and surface area of cones, spheres and pyramids",
  "Pythagoras' theorem",
  "Trigonometry in right-angled triangles (SOHCAHTOA)",
  "Angles of elevation and depression",
  "Bearings",
  "Pythagoras and trigonometry in 3D",
  "The sine rule and the cosine rule",
  "Area of a triangle using ½ab sin C",
  "Constructions and loci",
  "Scale drawings and maps",
  "Symmetry",
 ]),
 ("5  Vectors and transformation geometry", [
  "Reflection, rotation, translation and enlargement",
  "Negative and fractional scale factors",
  "Combined transformations",
  "Vector notation, magnitude, addition and scalar multiples",
  "Vector geometry and vector proof",
 ]),
 ("6  Statistics and probability", [
  "Collecting data: tally charts and frequency tables",
  "Two-way tables",
  "Bar charts, pictograms and pie charts",
  "Mean, median, mode and range",
  "Mean from a frequency table",
  "Estimated mean from grouped data",
  "Scatter graphs, correlation and the line of best fit",
  "Cumulative frequency graphs",
  "Box plots and quartiles",
  "Histograms with unequal class widths",
  "The probability scale and single-event probability",
  "Mutually exclusive and exhaustive events",
  "Relative frequency and expected frequency",
  "Sample space diagrams",
  "Tree diagrams, with replacement",
  "Tree diagrams without replacement; conditional probability",
  "Venn diagrams in probability",
 ]),
]

RATE_COL = 2  # column B

wb = Workbook()
ws = wb.active
ws.title = "Topic check"
ws.sheet_view.showGridLines = False

thin = Side(style="thin", color=HAIR)
box = Border(bottom=thin)


def put(sheet, r, c, v, **kw):
    cell = sheet.cell(row=r, column=c, value=v)
    cell.font = Font(name=F, size=kw.get("size", 11), bold=kw.get("bold", False),
                     color=kw.get("color", "000000"), italic=kw.get("italic", False))
    if kw.get("fill"):
        cell.fill = PatternFill("solid", fgColor=kw["fill"])
    cell.alignment = Alignment(horizontal=kw.get("h", "left"),
                               vertical="center", wrap_text=kw.get("wrap", False))
    if kw.get("border"):
        cell.border = box
    return cell


# ---------- header ----------
put(ws, 1, 1, "Edexcel IGCSE Mathematics A (4MA1) Higher Tier  —  Topic check",
    size=17, bold=True, color=NAVY)
put(ws, 2, 1, "Please fill this in before our first lesson. It takes about ten minutes. "
              "Don't think too hard about any single row.", size=10.5, color=SLATE)

for i, lab in enumerate(["NAME", "TARGET GRADE", "DATE"]):
    put(ws, 4, i + 1, lab, size=9, bold=True, color=AMBER_D)
    put(ws, 5, i + 1, "", fill=TINT, border=True)

for i, (name, meaning, bg, tx) in enumerate([
        ("Green", "I could do this in a test tomorrow", G_BG, G_TX),
        ("Amber", "I've seen it, but I'd need reminding", A_BG, A_TX),
        ("Red",   "I don't recognise this, or I always get it wrong", R_BG, R_TX)]):
    put(ws, 7 + i, 1, f"{name}  —  {meaning}", size=10.5, fill=bg, color=tx)

put(ws, 10, 1, "There is no wrong answer here — an honest red is worth far more to me "
               "than a hopeful green.", size=10, italic=True, color=SLATE)

HEAD = 12
FIRST = HEAD + 1
LAST = FIRST + sum(len(v) + 1 for _, v in TOPICS) - 1
RC = get_column_letter(RATE_COL)

put(ws, 7, 3, "RUNNING TOTAL", size=9, bold=True, color=AMBER_D)
ws.cell(row=8, column=3,
        value=f'=COUNTIF({RC}{FIRST}:{RC}{LAST},"Green")&" green    "&'
              f'COUNTIF({RC}{FIRST}:{RC}{LAST},"Amber")&" amber    "&'
              f'COUNTIF({RC}{FIRST}:{RC}{LAST},"Red")&" red"').font = Font(
    name=F, size=12, bold=True, color=NAVY)
put(ws, 9, 3, "Fills in as you go.", size=9.5, italic=True, color=SLATE)

# ---------- table ----------
for i, h in enumerate(["Topic", "How I feel about it", "Anything you want to tell me"], start=1):
    put(ws, HEAD, i, h, bold=True, color=WHITE, fill=NAVY, size=10.5,
        h="center" if i == RATE_COL else "left")
ws.row_dimensions[HEAD].height = 22

row = FIRST
area_ranges = []
for area, items in TOPICS:
    put(ws, row, 1, area, bold=True, color=NAVY, fill=TINT, size=11.5)
    for c in (2, 3):
        put(ws, row, c, "", fill=TINT)
    ws.row_dimensions[row].height = 21
    row += 1
    start = row
    for topic in items:
        put(ws, row, 1, topic, size=10.5, border=True)
        put(ws, row, RATE_COL, "", h="center", border=True)
        put(ws, row, 3, "", size=10, color=SLATE, border=True)
        ws.row_dimensions[row].height = 18
        row += 1
    area_ranges.append((area, start, row - 1))

# ---------- dropdown + colour, topic rows only ----------
dv = DataValidation(type="list", formula1='"Green,Amber,Red"',
                    allow_blank=True, showDropDown=False)
dv.promptTitle = "How do you feel about this topic?"
dv.prompt = "Green, Amber or Red"
ws.add_data_validation(dv)
topic_rngs = [f"{RC}{a}:{RC}{b}" for _, a, b in area_ranges]
for tr in topic_rngs:
    dv.add(tr)
for word, bg, tx in [("Green", G_BG, G_TX), ("Amber", A_BG, A_TX), ("Red", R_BG, R_TX)]:
    for tr in topic_rngs:
        ws.conditional_formatting.add(tr, CellIsRule(
            operator="equal", formula=[f'"{word}"'],
            fill=PatternFill("solid", start_color=bg, end_color=bg),
            font=Font(name=F, size=10.5, bold=True, color=tx)))

for col, w in zip("ABC", (82, 22, 48)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = f"A{FIRST}"
ws.print_title_rows = f"{HEAD}:{HEAD}"
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_area = f"A1:C{LAST}"

# ============ sheet 2: summary ============
s2 = wb.create_sheet("Summary")
s2.sheet_view.showGridLines = False
put(s2, 1, 1, "Summary", size=18, bold=True, color=NAVY)
put(s2, 2, 1, "Fills in on its own as the topic sheet is completed.", size=10.5, color=SLATE)
for i, h in enumerate(["Area", "Green", "Amber", "Red", "Not yet rated"], start=1):
    put(s2, 4, i, h, bold=True, color=WHITE, fill=NAVY, size=10.5,
        h="left" if i == 1 else "center")
s2.row_dimensions[4].height = 22

r = 5
for area, a, b in area_ranges:
    put(s2, r, 1, area, size=10.5, border=True)
    for j, (word, tx) in enumerate(zip(["Green", "Amber", "Red"], [G_TX, A_TX, R_TX]), start=2):
        put(s2, r, j, f'=COUNTIF(\'Topic check\'!{RC}{a}:{RC}{b},"{word}")',
            border=True, h="center", color=tx, bold=True, size=10.5)
    put(s2, r, 5, f"={b - a + 1}-SUM(B{r}:D{r})", border=True, h="center",
        color=SLATE, size=10.5)
    s2.row_dimensions[r].height = 19
    r += 1

put(s2, r, 1, "Total", bold=True, color=NAVY, fill=TINT, size=11)
for j in range(2, 6):
    L = get_column_letter(j)
    put(s2, r, j, f"=SUM({L}5:{L}{r-1})", bold=True, h="center", fill=TINT,
        size=11, color=NAVY)
s2.row_dimensions[r].height = 21
put(s2, r + 2, 1, "We start with the reds in the biggest areas. Ambers get folded into "
                  "practice. Greens we leave alone.", size=11, italic=True, color=NAVY)
for col, w in zip("ABCDE", (40, 12, 12, 12, 16)):
    s2.column_dimensions[col].width = w

# ============ sheet 3: practice ============
s3 = wb.create_sheet("Where to practise")
s3.sheet_view.showGridLines = False
put(s3, 1, 1, "Where to practise", size=18, bold=True, color=NAVY)
put(s3, 2, 1, "All free. Search any topic name from the checklist and you'll find "
              "questions on it.", size=10.5, color=SLATE)
for i, h in enumerate(["Site", "What it's best for"], start=1):
    put(s3, 4, i, h, bold=True, color=WHITE, fill=NAVY, size=10.5)
s3.row_dimensions[4].height = 22
sites = [
 ("Maths Genie", "Past-paper questions sorted by topic and by grade, with worked solutions. Start here."),
 ("Physics & Maths Tutor", "Edexcel IGCSE 4MA1 questions grouped by topic, with mark schemes."),
 ("Corbettmaths", "A short video and a worksheet for every topic. The 5-a-day sheets are ideal for daily practice."),
 ("Dr Frost Maths", "Free account, marks itself, and remembers what you got wrong."),
 ("Save My Exams", "Revision notes and topic questions for 4MA1 Higher specifically."),
 ("Past papers", "4MA1 Higher papers back to 2018, plus the R variants. Always use the real mark scheme."),
]
r = 5
for name, what in sites:
    put(s3, r, 1, name, size=11, bold=True, color=NAVY, border=True)
    put(s3, r, 2, what, size=10.5, color=SLATE, border=True)
    s3.row_dimensions[r].height = 20
    r += 1
put(s3, r + 1, 1, "Check you're getting IGCSE 4MA1 Higher material, not domestic GCSE — "
                  "the overlap is large but not total, and both IGCSE papers allow a calculator.",
    size=10.5, italic=True, color=NAVY)
s3.column_dimensions["A"].width = 26
s3.column_dimensions["B"].width = 86

OUT = ("/tmp/claude-0/-home-user-Daily/66a6df15-50ce-540f-aac0-2346f980a922/"
       "scratchpad/IGCSE-Maths-4MA1-Higher-Topic-Check.xlsx")
wb.save(OUT)
print("saved", OUT)
print("topics:", sum(len(v) for _, v in TOPICS), "| rating col", RC, "| rows", FIRST, "-", LAST)
