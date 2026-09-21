import json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
wb=load_workbook("Charlotte_RAG_mathsgenie.xlsx")
ws=wb["RAG tracker"]; last=ws.max_row
NAVY="FF162947"; TINT="FFEFEBE1"; RED="FFF4CCCC"; AMB="FFFCE5CD"; GRN="FFD9EAD3"
hdrF=Font(name="Arial",size=10,bold=True,color="FFF6F4EF"); hdrFill=PatternFill("solid",fgColor=NAVY)
bodyF=Font(name="Arial",size=10); thin=Side(style="thin",color="FFD8D3C6"); bd=Border(bottom=thin)
wrap=Alignment(wrap_text=True,vertical="top")

s3=wb.create_sheet("Summary",0)
s3["A1"]="Summary"; s3["A1"].font=Font(name="Arial",size=16,bold=True,color=NAVY)
s3["A2"]="Fills in automatically once the RAG tracker is completed."
s3["A2"].font=Font(name="Arial",size=10,italic=True,color="FF5D6471")
s3.append([]); s3.append(["Rating","Topics at grades 1 to 6","Rough workload","Questions each","Questions total"])
for c in s3[4]: c.font=hdrF; c.fill=hdrFill
rr=4
for val,act,q in [("R","Full worksheet, taught first",13),("A","Work until 3 right in a row",6),("G","2 question spot check",2)]:
    rr+=1
    s3.cell(rr,1,val); s3.cell(rr,2,f"=COUNTIFS('RAG tracker'!$C$2:$C${last},$A{rr},'RAG tracker'!$A$2:$A${last},\"<=6\")")
    s3.cell(rr,3,act); s3.cell(rr,4,q); s3.cell(rr,5,f"=$B{rr}*$D{rr}")
    for c in range(1,6): s3.cell(rr,c).font=bodyF; s3.cell(rr,c).border=bd
    s3.cell(rr,1).fill=PatternFill("solid",fgColor={"R":RED,"A":AMB,"G":GRN}[val])
rr+=1
s3.cell(rr,1,"Not yet rated"); s3.cell(rr,2,f"=COUNTIFS('RAG tracker'!$C$2:$C${last},\"\",'RAG tracker'!$A$2:$A${last},\">0\")")
s3.cell(rr,3,"Counts every row, all grades")
for c in range(1,6): s3.cell(rr,c).font=bodyF; s3.cell(rr,c).border=bd
rr+=1
s3.cell(rr,3,"Total questions"); s3.cell(rr,5,f"=SUM(E5:E7)")
for c in (3,5): s3.cell(rr,c).font=Font(name="Arial",size=10,bold=True); s3.cell(rr,c).fill=PatternFill("solid",fgColor=TINT)
rr+=1; s3.cell(rr,3,"Hours at about 2.5 min a question"); s3.cell(rr,5,f"=ROUND($E{rr-1}*2.5/60,1)")
rr+=1; s3.cell(rr,3,"Hours per week over six weeks"); s3.cell(rr,5,f"=ROUND($E{rr-1}/6,1)")
for r2 in (rr-1,rr):
    s3.cell(r2,3).font=bodyF; s3.cell(r2,5).font=bodyF
rr+=2
s3.cell(rr,1,"By grade").font=Font(name="Arial",size=11,bold=True,color=NAVY); rr+=1
for c,h in enumerate(["Grade","Topics","","Red","Amber","Green"],1): s3.cell(rr,c,h).font=hdrF; s3.cell(rr,c).fill=hdrFill
gstart=rr+1
for g in range(1,9):
    rr+=1
    s3.cell(rr,1,g)
    s3.cell(rr,2,f"=COUNTIFS('RAG tracker'!$A$2:$A${last},$A{rr})")
    for c,val in ((4,"R"),(5,"A"),(6,"G")):
        s3.cell(rr,c,f"=COUNTIFS('RAG tracker'!$A$2:$A${last},$A{rr},'RAG tracker'!$C$2:$C${last},\"{val}\")")
    for c in range(1,7): s3.cell(rr,c).font=bodyF; s3.cell(rr,c).border=bd
rr+=1
s3.cell(rr,1,"Total"); s3.cell(rr,2,f"=SUM(B{gstart}:B{rr-1})")
for c in (4,5,6): s3.cell(rr,c,f"=SUM({get_column_letter(c)}{gstart}:{get_column_letter(c)}{rr-1})")
for c in range(1,7): s3.cell(rr,c).font=Font(name="Arial",size=10,bold=True); s3.cell(rr,c).fill=PatternFill("solid",fgColor=TINT)
for col,wd in zip(range(1,7),[18,12,38,16,16,16]): s3.column_dimensions[get_column_letter(col)].width=wd

wb.save("Charlotte_RAG_mathsgenie.xlsx"); print("saved")
